from __future__ import annotations

import gzip
import hashlib
import io
import json
import os
import posixpath
import re
import time
import uuid
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

TZ = ZoneInfo("America/Sao_Paulo")
ROOT = Path(__file__).parent
RUNTIME = ROOT / ".runtime"
LOCAL_FILE = RUNTIME / "monitor_data.json"
UPLOAD_DIR = RUNTIME / "uploads"

SUPABASE_URL = os.getenv("SUPABASE_URL", "").strip()
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY", "").strip()

CENTRAL_SUPABASE_URL = "https://cuixazpxkvniqldmmnth.supabase.co"
CENTRAL_SUPABASE_ANON_KEY = (
    os.getenv("SETTA_SUPABASE_ANON_KEY", "").strip()
    or "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImN1aXhhenB4a3ZuaXFsZG1tbnRoIiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODc1MTYwNTMsImV4cCI6MjEwMzA5MjA1M30.jNFaIG1FcDYnMAoVaI23UYMuRL1BpZmuqu_LPEYb88E"
)
CENTRAL_EDGE_URL = f"{CENTRAL_SUPABASE_URL}/functions/v1/setta-data-api"

_client = None


def now_iso() -> str:
    return datetime.now(TZ).isoformat()


def parse_dt(value: Any) -> datetime | None:
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return dt.astimezone(TZ)
    except Exception:
        return None


def format_dt(value: Any) -> str:
    dt = parse_dt(value)
    return dt.strftime("%d/%m/%Y %H:%M:%S") if dt else "—"


def central_api_call(action: str, payload: dict | None = None, timeout: int = 60) -> dict:
    response = requests.post(
        CENTRAL_EDGE_URL,
        headers={
            "Authorization": f"Bearer {CENTRAL_SUPABASE_ANON_KEY}",
            "apikey": CENTRAL_SUPABASE_ANON_KEY,
            "Content-Type": "application/json",
        },
        json={"action": action, "payload": payload or {}},
        timeout=timeout,
    )
    try:
        data = response.json()
    except Exception:
        data = {"ok": False, "error": response.text or f"HTTP {response.status_code}"}
    if not response.ok or not data.get("ok"):
        raise RuntimeError(data.get("error") or f"HTTP {response.status_code}")
    return data


def central_client():
    from supabase import create_client
    return create_client(CENTRAL_SUPABASE_URL, CENTRAL_SUPABASE_ANON_KEY)


def supabase_enabled() -> bool:
    return bool(SUPABASE_URL and SUPABASE_KEY)


def monitor_shared_enabled() -> bool:
    # O backend compartilhado do Monitor vive na Central de Dados SETTA.
    # A Edge Function mantém o service_role somente no servidor.
    return bool(CENTRAL_SUPABASE_URL and CENTRAL_SUPABASE_ANON_KEY)


def get_client():
    global _client
    if not supabase_enabled():
        return None
    if _client is None:
        from supabase import create_client
        _client = create_client(SUPABASE_URL, SUPABASE_KEY)
    return _client


def _default_local() -> dict:
    return {
        "apis": [],
        "checks": [],
        "incidents": [],
        "sources": [],
        "imports": [],
        "visual_config": {},
    }


def _load_local() -> dict:
    RUNTIME.mkdir(parents=True, exist_ok=True)
    if not LOCAL_FILE.exists():
        return _default_local()
    try:
        data = json.loads(LOCAL_FILE.read_text(encoding="utf-8"))
        base = _default_local()
        base.update(data if isinstance(data, dict) else {})
        return base
    except Exception:
        return _default_local()


def _save_local(data: dict) -> None:
    RUNTIME.mkdir(parents=True, exist_ok=True)
    LOCAL_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def system_apis() -> list[dict]:
    central_endpoint = CENTRAL_EDGE_URL
    return [
        {
            "id": "11111111-1111-4111-8111-111111111101",
            "name": "CENTRAL DE DADOS SETTA",
            "app_name": "Saúde da Edge Function principal",
            "endpoint": central_endpoint + "/health",
            "method": "GET",
            "expected_status": 200,
            "timeout_seconds": 15,
            "warning_latency_ms": 1200,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "system": True,
            "source_app": "MONITOR DE APIs",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Supabase Edge Function",
            "resource": "setta-data-api /health",
            "code_location": "MONITOR-APIs/data_store.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111102",
            "name": "MRP → CENTRAL DE DADOS",
            "app_name": "Leitura das bases que alimentam o MRP",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 20,
            "warning_latency_ms": 1800,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_ACTION",
            "request_json": {
                "action": "bundle_state",
                "payload": {
                    "source_keys": ["cadastros"],
                    "derived_keys": [
                        "relatorio_geral_tratado",
                        "estoque_tratado",
                        "compras_tratado",
                        "tctp_tratado",
                        "relatorio_mrp",
                    ],
                },
            },
            "system": True,
            "source_app": "MRP",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Supabase Edge Function",
            "resource": "setta-data-api · bundle_state",
            "code_location": "MRP/central_mrp_data.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111103",
            "name": "CONVERSOR MRP → CENTRAL",
            "app_name": "Estado dos relatórios de origem e bases tratadas",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 20,
            "warning_latency_ms": 1800,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_PIPELINE_STATE",
            "request_json": {
                "action": "pipeline_state",
                "payload": {
                    "source_keys": ["relatorio_geral", "for001", "for022"],
                    "derived_key": "relatorio_geral_tratado",
                },
            },
            "system": True,
            "source_app": "CONVERSOR MRP",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Supabase Edge Function",
            "resource": "setta-data-api · pipeline_state",
            "code_location": "MRP-CONVERSOR/central_data.py · pipeline_sync.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111104",
            "name": "CENTRAL → STORAGE MRP",
            "app_name": "Disponibilização das bases tratadas",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 25,
            "warning_latency_ms": 2500,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_DERIVED_DOWNLOAD",
            "request_json": {
                "action": "derived_download",
                "payload": {"base_key": "relatorio_geral_tratado"},
            },
            "system": True,
            "source_app": "CENTRAL DE DADOS SETTA",
            "target_app": "SUPABASE STORAGE",
            "service": "Supabase Storage",
            "resource": "bucket setta-data · relatorio_geral_tratado",
            "code_location": "MRP/central_mrp_data.py · MRP-CONVERSOR/central_data.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111107",
            "name": "ESTOQUE TRATADO — PIPELINE",
            "app_name": "Analítico + Endereço → Estoque Tratado",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 20,
            "warning_latency_ms": 1800,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_PIPELINE_STATE",
            "request_json": {
                "action": "pipeline_state",
                "payload": {
                    "source_keys": ["analitico", "endereco"],
                    "derived_key": "estoque_tratado",
                },
            },
            "system": True,
            "source_app": "CONVERSOR MRP",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Edge Function + Storage",
            "resource": "pipeline estoque_tratado",
            "code_location": "MRP-CONVERSOR/conversores/estoque.py · pipeline_sync.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111108",
            "name": "COMPRAS TRATADO — PIPELINE",
            "app_name": "S.C + P.C + Pré-nota → Compras Tratado",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 20,
            "warning_latency_ms": 1800,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_PIPELINE_STATE",
            "request_json": {
                "action": "pipeline_state",
                "payload": {
                    "source_keys": ["sc", "pc", "pre_nota"],
                    "derived_key": "compras_tratado",
                },
            },
            "system": True,
            "source_app": "CONVERSOR MRP",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Edge Function + Storage",
            "resource": "pipeline compras_tratado",
            "code_location": "MRP-CONVERSOR/conversores/compras.py · pipeline_sync.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111109",
            "name": "TCTP TRATADO — PIPELINE",
            "app_name": "PMP + H001 → TCTP Tratado",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 20,
            "warning_latency_ms": 1800,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_PIPELINE_STATE",
            "request_json": {
                "action": "pipeline_state",
                "payload": {
                    "source_keys": ["pmp", "h001"],
                    "derived_key": "tctp_tratado",
                },
            },
            "system": True,
            "source_app": "CONVERSOR MRP",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Edge Function + Storage",
            "resource": "pipeline tctp_tratado",
            "code_location": "MRP-CONVERSOR/conversores/tc_tp.py · pipeline_sync.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111110",
            "name": "INVENTÁRIO ROTATIVO → CENTRAL",
            "app_name": "Sincronização de Analítico e Endereço",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 20,
            "warning_latency_ms": 1800,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_CONSUMER_SYNC",
            "expected_source_keys": ["analitico", "endereco"],
            "request_json": {
                "action": "consumer_sync_status",
                "payload": {"consumer_key": "inventario_rotativo"},
            },
            "system": True,
            "source_app": "INVENTÁRIO ROTATIVO",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Supabase Edge Function",
            "resource": "setta-data-api · consumer_sync_status",
            "code_location": "INVENTARIO-ROTATIVO-SETTA/central_inventory_data.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111111",
            "name": "FECHAMENTO MENSAL → CENTRAL",
            "app_name": "Sincronização de Analítico e Cadastros",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 20,
            "warning_latency_ms": 1800,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_CONSUMER_SYNC",
            "expected_source_keys": ["analitico", "cadastros"],
            "request_json": {
                "action": "consumer_sync_status",
                "payload": {"consumer_key": "fechamento_mensal"},
            },
            "system": True,
            "source_app": "FECHAMENTO MENSAL",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Supabase Edge Function",
            "resource": "setta-data-api · consumer_sync_status",
            "code_location": "FECHAMENTO-MENSAL/central_fechamento_data.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111112",
            "name": "CONTROLE DE NFS → CENTRAL",
            "app_name": "Sincronização de NF e MES Pré-notas",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 20,
            "warning_latency_ms": 1800,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_CONSUMER_SYNC",
            "expected_source_keys": ["nf", "mes_pre_notas"],
            "request_json": {
                "action": "consumer_sync_status",
                "payload": {"consumer_key": "controle_nfs"},
            },
            "system": True,
            "source_app": "CONTROLE DE NFS",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Supabase Edge Function",
            "resource": "setta-data-api · consumer_sync_status",
            "code_location": "CONTROLE-DE-NFS/central_nfs_data.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111113",
            "name": "ALMOXARIFADO — CONVITE DE USUÁRIO",
            "app_name": "Disponibilidade da função de administração de usuários",
            "endpoint": (
                "https://cuixazpxkvniqldmmnth.supabase.co/"
                "functions/v1/almox-convidar-usuario"
            ),
            "method": "GET",
            "expected_status": 405,
            "timeout_seconds": 15,
            "warning_latency_ms": 1200,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "system": True,
            "source_app": "GESTÃO DE ESTOQUE",
            "target_app": "SUPABASE AUTH",
            "service": "Supabase Edge Function",
            "resource": "almox-convidar-usuario",
            "code_location": "Supabase Edge Function · almox-convidar-usuario",
        },
        {
            "id": "11111111-1111-4111-8111-111111111114",
            "name": "MRP → HISTÓRICO",
            "app_name": "Persistência dos snapshots do MRP",
            "endpoint": (
                "https://cuixazpxkvniqldmmnth.supabase.co/"
                "rest/v1/mrp_snapshots?select=*&limit=1"
            ),
            "method": "GET",
            "expected_status": 200,
            "timeout_seconds": 15,
            "warning_latency_ms": 1200,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "system": True,
            "source_app": "MRP",
            "target_app": "SUPABASE DATABASE",
            "service": "PostgREST / PostgreSQL",
            "resource": "mrp_snapshots",
            "code_location": "MRP/mrp_conexao_patch.py · mrp_salvamento_compacto_patch.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111115",
            "name": "MRP → TRATATIVAS DE PROJETO",
            "app_name": "Persistência das tratativas por projeto",
            "endpoint": (
                "https://cuixazpxkvniqldmmnth.supabase.co/"
                "rest/v1/mrp_project_treatments?select=*&limit=1"
            ),
            "method": "GET",
            "expected_status": 200,
            "timeout_seconds": 15,
            "warning_latency_ms": 1200,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "system": True,
            "source_app": "MRP",
            "target_app": "SUPABASE DATABASE",
            "service": "PostgREST / PostgreSQL",
            "resource": "mrp_project_treatments",
            "code_location": "MRP/app_mrp_runtime.py · mrp_tratativa_produto_patch.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111117",
            "name": "MRP → TRATATIVAS DE PRODUTO",
            "app_name": "Persistência das tratativas por produto",
            "endpoint": (
                "https://cuixazpxkvniqldmmnth.supabase.co/"
                "rest/v1/mrp_project_product_treatments?select=*&limit=1"
            ),
            "method": "GET",
            "expected_status": 200,
            "timeout_seconds": 15,
            "warning_latency_ms": 1200,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "system": True,
            "source_app": "MRP",
            "target_app": "SUPABASE DATABASE",
            "service": "PostgREST / PostgreSQL",
            "resource": "mrp_project_product_treatments",
            "code_location": "MRP/mrp_tratativa_produto_patch.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111118",
            "name": "INVENTÁRIO → ESTOQUE TRATADO",
            "app_name": "Leitura da base tratada no Banco de Dados do Inventário",
            "endpoint": central_endpoint,
            "method": "POST",
            "expected_status": 200,
            "timeout_seconds": 25,
            "warning_latency_ms": 2500,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "CENTRAL_DERIVED_DOWNLOAD",
            "request_json": {
                "action": "derived_download",
                "payload": {"base_key": "estoque_tratado"},
            },
            "system": True,
            "source_app": "INVENTÁRIO ROTATIVO",
            "target_app": "SUPABASE STORAGE",
            "service": "Supabase Edge Function + Storage",
            "resource": "setta-data · estoque_tratado",
            "code_location": "INVENTARIO-ROTATIVO-SETTA/central_inventory_data.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111116",
            "name": "SETTA HUB → CONFIGURAÇÃO",
            "app_name": "Persistência das configurações do Hub",
            "endpoint": (
                "https://cuixazpxkvniqldmmnth.supabase.co/"
                "rest/v1/setta_hub_config?select=*&limit=1"
            ),
            "method": "GET",
            "expected_status": 200,
            "timeout_seconds": 15,
            "warning_latency_ms": 1200,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "system": True,
            "source_app": "SETTA HUB",
            "target_app": "SUPABASE DATABASE",
            "service": "PostgREST",
            "resource": "setta_hub_config",
            "code_location": "SETTA-HUB/streamlit_app.py",
        },
        {
            "id": "11111111-1111-4111-8111-111111111105",
            "name": "GESTÃO DE ENTREGAS — OPERACIONAL",
            "app_name": "Cronograma • Central • Sincronização",
            "endpoint": (
                "https://cuixazpxkvniqldmmnth.supabase.co/"
                "functions/v1/entrega-cronograma-api/health"
            ),
            "method": "GET",
            "expected_status": 200,
            "timeout_seconds": 15,
            "warning_latency_ms": 1200,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "system": True,
            "source_app": "GESTÃO DE ENTREGAS",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Supabase Edge Function",
            "resource": "entrega-cronograma-api /health",
            "code_location": "Gestão de Entregas · integração operacional",
        },
        {
            "id": "11111111-1111-4111-8111-111111111106",
            "name": "GESTÃO DE ENTREGAS → NFs",
            "app_name": "Materiais pendentes • Impacto MRP",
            "endpoint": (
                "https://cuixazpxkvniqldmmnth.supabase.co/"
                "functions/v1/nf-materiais-api/status"
            ),
            "method": "GET",
            "expected_status": 200,
            "timeout_seconds": 15,
            "warning_latency_ms": 1500,
            "active": True,
            "secret_ref": "__GESTAO_ENTREGAS_ANON__",
            "health_mode": "NF_MATERIAIS_STATUS",
            "system": True,
            "source_app": "GESTÃO DE ENTREGAS",
            "target_app": "CONTROLE DE NFs",
            "service": "Supabase Edge Function",
            "resource": "nf-materiais-api /status",
            "code_location": "Gestão de Entregas · Controle de NFs",
        },
        {
            "id": "11111111-1111-4111-8111-111111111119",
            "name": "ERP INDUSTRIAL → PMP",
            "app_name": "Recebimento autenticado do PMP para o TC/TP Tratado",
            "endpoint": "https://cuixazpxkvniqldmmnth.supabase.co/functions/v1/pmp-ingest/health",
            "method": "GET",
            "expected_status": 200,
            "timeout_seconds": 15,
            "warning_latency_ms": 1500,
            "active": True,
            "secret_ref": "",
            "system": True,
            "source_app": "ERP INDUSTRIAL",
            "target_app": "CENTRAL DE DADOS SETTA",
            "service": "Supabase Edge Function",
            "resource": "pmp-ingest /health (token configurado)",
            "code_location": "Supabase Edge Function · pmp-ingest",
        },
    ]


def _system_db_payload(api: dict) -> dict:
    return {
        "id": str(api["id"]),
        "name": str(api.get("name") or ""),
        "app_name": str(api.get("app_name") or ""),
        "endpoint": str(api.get("endpoint") or ""),
        "method": str(api.get("method") or "GET"),
        "expected_status": int(api.get("expected_status") or 200),
        "timeout_seconds": int(api.get("timeout_seconds") or 10),
        "warning_latency_ms": int(api.get("warning_latency_ms") or 1000),
        "active": bool(api.get("active", True)),
        "secret_ref": str(api.get("secret_ref") or ""),
    }


def ensure_system_apis() -> None:
    global _system_sync_done
    if globals().get("_system_sync_done"):
        return
    rows = [_system_db_payload(item) for item in system_apis()]
    try:
        central_api_call("monitor_apis_upsert", {"rows": rows}, timeout=30)
        _system_sync_done = True
    except Exception:
        # Não usar armazenamento local silenciosamente: a UI deve refletir
        # a indisponibilidade do backend compartilhado.
        _system_sync_done = False



def connection_coverage() -> list[dict]:
    monitored = [
        {
            "connection": item["name"],
            "source": item.get("source_app") or "—",
            "target": item.get("target_app") or "—",
            "service": item.get("service") or "—",
            "resource": item.get("resource") or "—",
            "code": item.get("code_location") or "—",
            "monitoring": "MONITORADO",
        }
        for item in system_apis()
    ]
    infrastructure = [
        {
            "connection": "MONITOR WORKER → MONITORAMENTO",
            "source": "SETTA MONITOR WORKER",
            "target": "API_CHECKS / INCIDENTES",
            "service": "Supabase Edge Function",
            "resource": "setta-monitor-worker",
            "code": "Supabase Edge Function · setta-monitor-worker",
            "monitoring": "AUTO-MONITORAMENTO INDIRETO",
        },
    ]
    return monitored + infrastructure


def _integration_location(item: dict) -> dict:
    source = str(item.get("source_app") or "").upper()
    target = str(item.get("target_app") or "").upper()
    name = str(item.get("name") or "").upper()
    haystack = " ".join((source, target, name))
    if "CONVERSOR MRP" in haystack:
        return {
            "service": "Edge Function + Storage",
            "resource": "setta-data-api · fontes e bases tratadas",
            "code_location": "MRP-CONVERSOR/central_data.py · pipeline_sync.py",
        }
    if "MRP" in haystack and "CENTRAL" in haystack:
        return {
            "service": "Edge Function + Storage",
            "resource": "setta-data-api · bundle/derived",
            "code_location": "MRP/central_mrp_data.py",
        }
    if "MRP" in haystack:
        return {
            "service": "Central de Dados SETTA",
            "resource": "relatorio_mrp",
            "code_location": "MRP/mrp_central_data_patch.py",
        }
    return {
        "service": str(item.get("service") or "Central de Dados SETTA"),
        "resource": str(item.get("resource") or "—"),
        "code_location": str(item.get("code_location") or "—"),
    }



def list_integrations() -> list[dict]:
    try:
        rows = central_api_call("integration_status", timeout=30).get("data") or []
    except Exception as exc:
        return [
            {
                "id": "integration-status-error",
                "name": "INTEGRAÇÕES SETTA",
                "source_app": "MONITOR DE APIS",
                "target_app": "CENTRAL DE DADOS",
                "status": "OFFLINE",
                "last_activity_at": None,
                "last_activity_label": "—",
                "resource_count": 0,
                "expected_count": 0,
                "rows_count": 0,
                "version_label": "FALHA AO CONSULTAR INTEGRAÇÕES",
                "resources": [],
                "error_message": str(exc),
            }
        ]

    output: list[dict] = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        item = dict(row)
        item["status"] = str(item.get("status") or "SEM DADOS").upper()
        item["last_activity_label"] = format_dt(item.get("last_activity_at"))
        item["resource_count"] = int(item.get("resource_count") or 0)
        item["expected_count"] = int(item.get("expected_count") or 0)
        item["rows_count"] = int(item.get("rows_count") or 0)
        item.update(_integration_location(item))
        output.append(item)
    return output


def list_apis(include_demo: bool = True) -> list[dict]:
    ensure_system_apis()
    try:
        rows = central_api_call("monitor_apis_list", timeout=30).get("data") or []
    except Exception:
        rows = []

    configs = system_apis()
    by_id = {str(row.get("id")): row for row in rows if isinstance(row, dict)}
    system_ids = {str(item.get("id")) for item in configs}
    systems = [
        {**config, **by_id.get(str(config["id"]), {})}
        for config in configs
    ]
    systems = [
        {
            **item,
            **{
                key: config.get(key)
                for key in (
                    "system", "health_mode", "request_json", "expected_source_keys",
                    "source_app", "target_app", "service", "resource", "code_location",
                )
            },
        }
        for item, config in zip(systems, configs)
    ]
    custom = [
        row for row in rows
        if isinstance(row, dict) and str(row.get("id")) not in system_ids
    ]
    return systems + custom



def save_api(payload: dict) -> dict:
    clean = {
        "id": str(payload.get("id") or "").strip() or None,
        "name": str(payload.get("name") or "").strip(),
        "app_name": str(payload.get("app_name") or "").strip(),
        "endpoint": str(payload.get("endpoint") or "").strip(),
        "method": str(payload.get("method") or "GET").upper(),
        "expected_status": int(payload.get("expected_status") or 200),
        "timeout_seconds": int(payload.get("timeout_seconds") or 10),
        "warning_latency_ms": int(payload.get("warning_latency_ms") or 1000),
        "active": bool(payload.get("active", True)),
        "secret_ref": str(payload.get("secret_ref") or "").strip(),
    }
    if not clean["name"] or not clean["endpoint"]:
        raise ValueError("Nome e endpoint são obrigatórios.")
    if clean["id"] is None:
        clean.pop("id")
    row = central_api_call(
        "monitor_api_save",
        clean,
        timeout=30,
    ).get("data") or clean
    return row



def delete_api(api_id: str) -> None:
    system_ids = {str(item.get("id")) for item in system_apis()}
    if str(api_id) in system_ids:
        return
    central_api_call(
        "monitor_api_delete",
        {"id": str(api_id)},
        timeout=30,
    )



def list_checks(api_id: str | None = None, limit: int = 200) -> list[dict]:
    try:
        return central_api_call(
            "monitor_checks_list",
            {"api_id": api_id or "", "limit": int(limit)},
            timeout=30,
        ).get("data") or []
    except Exception:
        return []


def list_check_summary(
    api_id: str,
    since_iso: str,
    limit: int = 2000,
) -> list[dict]:
    grouped = list_check_summaries([api_id], since_iso, limit)
    return grouped.get(str(api_id), [])


def list_check_summaries(
    api_ids: list[str],
    since_iso: str,
    limit: int = 20000,
) -> dict[str, list[dict]]:
    ids = [str(value) for value in api_ids if str(value or "").strip()]
    grouped = {api_id: [] for api_id in ids}
    if not ids:
        return grouped
    try:
        rows = central_api_call(
            "monitor_checks_summary",
            {
                "api_ids": ids,
                "since_iso": str(since_iso or ""),
                "limit": int(limit),
            },
            timeout=30,
        ).get("data") or []
    except Exception:
        return grouped
    for row in rows:
        if not isinstance(row, dict):
            continue
        api_id = str(row.get("api_id") or "")
        if api_id in grouped:
            grouped[api_id].append(row)
    return grouped



def save_check(api_id: str, result: dict) -> dict:
    row = {
        "api_id": str(api_id),
        "checked_at": result.get("checked_at") or now_iso(),
        "status": result.get("status"),
        "http_status": result.get("http_status"),
        "latency_ms": result.get("latency_ms"),
        "success": bool(result.get("success")),
        "error_message": str(result.get("error_message") or "")[:1000],
    }
    return (
        central_api_call(
            "monitor_check_save",
            {"row": row},
            timeout=30,
        ).get("data")
        or row
    )



def update_api_runtime(api_id: str, fields: dict) -> None:
    if str(api_id).startswith("demo-"):
        return
    allowed = {
        "status", "last_check_at", "last_success_at", "last_failure_at",
        "last_http_status", "last_latency_ms", "consecutive_failures",
    }
    clean = {key: value for key, value in fields.items() if key in allowed}
    central_api_call(
        "monitor_runtime_update",
        {"id": str(api_id), "fields": clean},
        timeout=30,
    )



def list_incidents(limit: int = 100) -> list[dict]:
    try:
        return central_api_call(
            "monitor_incidents_list",
            {"limit": int(limit)},
            timeout=30,
        ).get("data") or []
    except Exception:
        return []


def open_incident(api_id: str, api_name: str, kind: str, message: str) -> None:
    central_api_call(
        "monitor_incident_open",
        {
            "api_id": str(api_id),
            "api_name": str(api_name),
            "kind": str(kind),
            "message": str(message)[:1000],
        },
        timeout=30,
    )


def resolve_incidents(api_id: str) -> None:
    central_api_call(
        "monitor_incidents_resolve",
        {"api_id": str(api_id)},
        timeout=30,
    )



SOURCE_CATALOG = [
    {
        "key": "relatorio_geral",
        "name": "RELATÓRIO GERAL",
        "source_system": "MES",
        "apps": "Conversor MRP • Gestão de Equipes",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "for001",
        "name": "FOR001",
        "source_system": "PLANILHA MANUAL",
        "apps": "Conversor MRP",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "for022",
        "name": "FOR022",
        "source_system": "PLANILHA MANUAL",
        "apps": "Conversor MRP • Gestão de Entregas • Gestão de Equipes",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "analitico",
        "name": "ANALÍTICO",
        "source_system": "PROTHEUS",
        "apps": "Conversor MRP • Fechamento Mensal • Inventário Rotativo",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "endereco",
        "name": "ENDEREÇO",
        "source_system": "PROTHEUS",
        "apps": "Conversor MRP • Inventário Rotativo",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "sc",
        "name": "S.C",
        "source_system": "PROTHEUS",
        "apps": "Conversor MRP",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "pc",
        "name": "P.C",
        "source_system": "PROTHEUS",
        "apps": "Conversor MRP",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "pre_nota",
        "name": "PRÉ NOTA",
        "source_system": "PROTHEUS",
        "apps": "Conversor MRP",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "pmp",
        "name": "PMP",
        "source_system": "OUTRO APP",
        "apps": "Conversor MRP",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "h001",
        "name": "H001",
        "source_system": "PLANILHA MANUAL",
        "apps": "Conversor MRP",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "cadastros",
        "name": "CADASTROS",
        "source_system": "PROTHEUS",
        "apps": "MRP • Gestão de Equipes",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "nf",
        "name": "NF",
        "source_system": "PROTHEUS",
        "apps": "Gestão de Entregas • Controle de NFs",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "mes_pre_notas",
        "name": "MES PRÉ NOTAS",
        "source_system": "MES",
        "apps": "Controle de NFs",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
    {
        "key": "movimentacao",
        "name": "MOVIMENTAÇÃO",
        "source_system": "PROTHEUS",
        "apps": "Fechamento Mensal • Gestão de Equipes",
        "mode": "UPLOAD CENTRAL",
        "api_plan": True,
    },
]

DERIVED_CATALOG = [
    {
        "key": "relatorio_geral_tratado",
        "name": "RELATÓRIO GERAL TRATADO",
        "producer": "Conversor MRP",
        "depends_on": ["RELATÓRIO GERAL", "FOR001", "FOR022"],
        "apps": "MRP",
        "mode": "AUTOMÁTICO",
    },
    {
        "key": "estoque_tratado",
        "name": "ESTOQUE TRATADO",
        "producer": "Conversor MRP",
        "depends_on": ["ANALÍTICO", "ENDEREÇO"],
        "apps": "MRP • Inventário Rotativo",
        "mode": "AUTOMÁTICO",
    },
    {
        "key": "compras_tratado",
        "name": "COMPRAS TRATADO",
        "producer": "Conversor MRP",
        "depends_on": ["S.C", "P.C", "PRÉ NOTA"],
        "apps": "MRP",
        "mode": "AUTOMÁTICO",
    },
    {
        "key": "tctp_tratado",
        "name": "TCTP TRATADO",
        "producer": "Conversor MRP",
        "depends_on": ["PMP", "H001"],
        "apps": "MRP",
        "mode": "AUTOMÁTICO",
    },
    {
        "key": "relatorio_mrp",
        "name": "RELATÓRIO MRP",
        "producer": "MRP",
        "depends_on": [
            "CADASTROS",
            "RELATÓRIO GERAL TRATADO",
            "ESTOQUE TRATADO",
            "COMPRAS TRATADO",
            "TCTP TRATADO",
        ],
        "apps": "Gestão de Entregas • Gestão de Equipes",
        "mode": "AUTOMÁTICO",
    },
    {
        "key": "materiais_api",
        "name": "MATERIAIS",
        "producer": "Gestão de Entregas",
        "depends_on": ["FOR022", "RELATÓRIO MRP", "NF"],
        "apps": "Controle de NFs",
        "mode": "API ATIVA",
    },
]



def source_operational_status(
    source_row: dict,
    normalized_row: dict,
) -> str:
    """Status diário real da fonte.

    A regra é calendário local (America/Sao_Paulo): ao virar 00:00,
    uma fonte não conferida no novo dia passa automaticamente para
    DESATUALIZADO, sem depender de job ou gravação no banco.
    """
    raw_status = str(source_row.get("status") or "").strip().upper()
    normalized_status = str(normalized_row.get("status") or "").strip().upper()

    if raw_status == "ERRO" or normalized_status == "ERRO":
        return "ERRO"

    if not bool(source_row.get("available")):
        return "PENDENTE"

    source_version = int(source_row.get("version") or 0)
    normalized_version = int(normalized_row.get("source_version") or 0)
    normalized_current = (
        bool(normalized_row.get("available"))
        and not bool(normalized_row.get("stale"))
        and source_version > 0
        and normalized_version == source_version
    )
    if not normalized_current:
        return "PENDENTE"

    today = datetime.now(TZ).date()
    received_at = parse_dt(
        source_row.get("last_received_at")
        or source_row.get("last_update_at")
    )
    if received_at is None or received_at.date() < today:
        return "DESATUALIZADO"

    content_at = parse_dt(source_row.get("last_update_at"))
    if content_at is not None and content_at.date() == today:
        return "ATUALIZADO"

    return "SEM ALTERAÇÃO"



def list_sources() -> list[dict]:
    keys = [item["key"] for item in SOURCE_CATALOG]
    try:
        rows = central_api_call(
            "source_status",
            {"keys": keys},
            timeout=30,
        ).get("data") or []
        by_key = {
            str(row.get("source_key")): row
            for row in rows
            if isinstance(row, dict)
        }
    except Exception:
        by_key = {}

    try:
        normalized_rows = central_api_call(
            "source_normalized_status",
            {"keys": keys},
            timeout=30,
        ).get("data") or []
        normalized_by_key = {
            str(row.get("source_key")): row
            for row in normalized_rows
            if isinstance(row, dict)
        }
    except Exception:
        normalized_by_key = {}

    result = []
    for base in SOURCE_CATALOG:
        row = by_key.get(base["key"], {})
        normalized = normalized_by_key.get(base["key"], {})
        normalized_current = bool(normalized.get("available"))
        source_payload = {
            **row,
            "available": bool(row.get("available")),
        }
        normalized_payload = {
            **normalized,
            "available": normalized_current,
        }
        operational_status = source_operational_status(
            source_payload,
            normalized_payload,
        )
        result.append(
            {
                "source_key": base["key"],
                "name": base["name"],
                "apps": base["apps"],
                "source_system": base["source_system"],
                "mode": base["mode"],
                "api_plan": bool(base.get("api_plan")),
                "status": operational_status,
                "stored_status": row.get("status") or "AGUARDANDO",
                "last_update_at": row.get("last_update_at"),
                "last_received_at": row.get("last_received_at") or row.get("last_update_at"),
                "rows_count": row.get("rows_count") or 0,
                "origin": row.get("origin") or base["source_system"],
                "last_file_name": row.get("last_file_name") or "",
                "content_sha256": row.get("content_sha256") or "",
                "version": int(row.get("version") or 0),
                "available": bool(row.get("available")),
                "normalized_available": normalized_current,
                "normalized_stale": bool(normalized.get("stale")),
                "normalized_version": int(normalized.get("source_version") or 0),
                "normalized_format": normalized.get("format") or "",
                "normalized_at": normalized.get("normalized_at"),
                "normalized_storage_path": (
                    normalized.get("storage_path") or ""
                    if normalized_current
                    else ""
                ),
            }
        )
    return result


def list_derived_bases() -> list[dict]:
    """Retorna bases derivadas com status por consistência de dependências.

    Uma base continua ATUALIZADA se as versões que a geraram ainda são as
    versões correntes das fontes. Não é necessário reprocessar só porque
    virou o dia quando nenhuma dependência mudou.
    """
    rows_by_key: dict[str, dict] = {}
    source_versions: dict[str, int] = {}

    try:
        rows = central_api_call(
            "derived_status",
            {"keys": [item["key"] for item in DERIVED_CATALOG]},
            timeout=30,
        ).get("data") or []
        rows_by_key = {
            str(row.get("base_key")): row
            for row in rows
            if isinstance(row, dict)
        }
    except Exception:
        rows_by_key = {}

    try:
        source_rows = central_api_call(
            "source_status",
            {"keys": [item["key"] for item in SOURCE_CATALOG]},
            timeout=30,
        ).get("data") or []
        source_versions = {
            str(row.get("source_key")): int(row.get("version") or 0)
            for row in source_rows
            if isinstance(row, dict)
        }
    except Exception:
        source_versions = {}

    def same_source_versions(row: dict, expected_keys: list[str]) -> bool:
        used = row.get("source_versions") or {}
        if not isinstance(used, dict):
            return False
        for key in expected_keys:
            try:
                if int(used.get(key) or 0) != int(source_versions.get(key) or 0):
                    return False
            except Exception:
                return False
        return True

    result = []
    for base in DERIVED_CATALOG:
        row = rows_by_key.get(base["key"], {})
        processed_at = row.get("processed_at")

        if base["key"] == "relatorio_geral_tratado":
            current = same_source_versions(
                row,
                ["relatorio_geral", "for001", "for022"],
            )
        elif base["key"] == "estoque_tratado":
            current = same_source_versions(
                row,
                ["analitico", "endereco"],
            )
        elif base["key"] == "compras_tratado":
            current = same_source_versions(
                row,
                ["sc", "pc", "pre_nota"],
            )
        elif base["key"] == "tctp_tratado":
            current = same_source_versions(
                row,
                ["pmp", "h001"],
            )
        elif base["key"] == "relatorio_mrp":
            used = row.get("source_versions") or {}
            cad_ok = str(used.get("cadastros") or "") == f"v{int(source_versions.get('cadastros') or 0)}"
            derived_ok = True
            for dep_key in [
                "relatorio_geral_tratado",
                "estoque_tratado",
                "compras_tratado",
                "tctp_tratado",
            ]:
                dep = rows_by_key.get(dep_key, {})
                dep_processed = str(dep.get("processed_at") or "")
                if str(used.get(dep_key) or "") != dep_processed:
                    derived_ok = False
                    break
            current = cad_ok and derived_ok
        else:
            # MATERIAIS é integração por API e não usa setta_derived_bases.
            current = bool(row.get("available"))

        if not bool(row.get("available")):
            operational_status = "AGUARDANDO"
        elif not current:
            operational_status = "DESATUALIZADO"
        else:
            operational_status = "ATUALIZADO"

        result.append(
            {
                **base,
                "status": operational_status,
                "stored_status": row.get("status") or "AGUARDANDO",
                "rows_count": int(row.get("rows_count") or 0),
                "processed_at": processed_at,
                "processed_label": format_dt(processed_at),
                "available": bool(row.get("available")),
                "source_versions": row.get("source_versions") or {},
            }
        )
    return result



def _frame_values(frame: pd.DataFrame) -> list[list[Any]]:
    """Converte DataFrame para valores JSON preservando datas e nulos."""
    if frame is None:
        return []
    text = frame.to_json(
        orient="values",
        date_format="iso",
        force_ascii=False,
    )
    return json.loads(text)


def _trim_workbook_row(values: tuple[Any, ...], max_columns: int = 256) -> list[Any]:
    """Mantém somente a parte útil da linha e limita used-range artificial."""
    row = list(values[:max_columns])
    while row and (
        row[-1] is None
        or (isinstance(row[-1], str) and not row[-1].strip())
    ):
        row.pop()
    return row


def _strip_invalid_conditional_formatting(raw: bytes) -> bytes:
    """Remove regras visuais de formatação condicional do XLSX/XLT(X).

    Alguns relatórios do ERP trazem referências de formatação condicional
    inválidas para versões recentes do openpyxl. A Central precisa somente
    dos valores das células, portanto essas regras podem ser descartadas
    na cópia usada exclusivamente durante a normalização.
    """
    source = io.BytesIO(raw)
    target = io.BytesIO()

    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(
        target,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if (
                info.filename.startswith("xl/worksheets/")
                and info.filename.endswith(".xml")
                and b"conditionalFormatting" in data
            ):
                data = re.sub(
                    rb"<(?:[A-Za-z0-9_]+:)?conditionalFormatting[^>]*>.*?</(?:[A-Za-z0-9_]+:)?conditionalFormatting>",
                    b"",
                    data,
                    flags=re.DOTALL,
                )
            zout.writestr(info, data)

    return target.getvalue()


def _load_workbook_streaming(raw: bytes):
    """Abre o workbook com contingência para MultiCellRange inválido."""
    try:
        return load_workbook(
            io.BytesIO(raw),
            read_only=True,
            data_only=True,
        )
    except TypeError as exc:
        if "MultiCellRange" not in str(exc):
            raise
        sanitized = _strip_invalid_conditional_formatting(raw)
        return load_workbook(
            io.BytesIO(sanitized),
            read_only=True,
            data_only=True,
        )


def _xlsx_column_index(cell_ref: str) -> int:
    letters = re.match(r"([A-Za-z]+)", str(cell_ref or ""))
    if not letters:
        return 0
    value = 0
    for char in letters.group(1).upper():
        value = value * 26 + (ord(char) - 64)
    return value


def _xlsx_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    path = "xl/sharedStrings.xml"
    if path not in archive.namelist():
        return []

    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    root = ET.fromstring(archive.read(path))
    values: list[str] = []
    for item in root.findall(f"{ns}si"):
        values.append(
            "".join(
                node.text or ""
                for node in item.iter(f"{ns}t")
            )
        )
    return values


def _xlsx_date_styles(archive: zipfile.ZipFile) -> set[int]:
    path = "xl/styles.xml"
    if path not in archive.namelist():
        return set()

    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    root = ET.fromstring(archive.read(path))

    custom_formats: dict[int, str] = {}
    num_fmts = root.find(f"{ns}numFmts")
    if num_fmts is not None:
        for item in num_fmts.findall(f"{ns}numFmt"):
            try:
                custom_formats[int(item.attrib.get("numFmtId", "0"))] = str(
                    item.attrib.get("formatCode") or ""
                )
            except Exception:
                continue

    date_styles: set[int] = set()
    cell_xfs = root.find(f"{ns}cellXfs")
    if cell_xfs is None:
        return date_styles

    for index, xf in enumerate(cell_xfs.findall(f"{ns}xf")):
        try:
            num_fmt_id = int(xf.attrib.get("numFmtId", "0"))
        except Exception:
            num_fmt_id = 0
        format_code = custom_formats.get(
            num_fmt_id,
            BUILTIN_FORMATS.get(num_fmt_id, ""),
        )
        try:
            if format_code and is_date_format(format_code):
                date_styles.add(index)
        except Exception:
            continue
    return date_styles


def _xlsx_workbook_manifest(
    archive: zipfile.ZipFile,
) -> tuple[list[tuple[str, str]], bool]:
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    rid_attr = (
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )

    root = ET.fromstring(archive.read("xl/workbook.xml"))
    workbook_pr = root.find(f"{ns}workbookPr")
    date1904 = bool(
        workbook_pr is not None
        and str(workbook_pr.attrib.get("date1904") or "").lower()
        in {"1", "true"}
    )

    rels_root = ET.fromstring(
        archive.read("xl/_rels/workbook.xml.rels")
    )
    relationships: dict[str, str] = {}
    for rel in rels_root.findall(f"{rel_ns}Relationship"):
        rel_id = str(rel.attrib.get("Id") or "")
        target = str(rel.attrib.get("Target") or "")
        if not rel_id or not target:
            continue
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = posixpath.normpath(posixpath.join("xl", target))
        relationships[rel_id] = target

    output: list[tuple[str, str]] = []
    sheets = root.find(f"{ns}sheets")
    if sheets is None:
        return output, date1904

    for sheet in sheets.findall(f"{ns}sheet"):
        name = str(sheet.attrib.get("name") or "Planilha")
        rel_id = str(sheet.attrib.get(rid_attr) or "")
        path = relationships.get(rel_id)
        if path and path in archive.namelist():
            output.append((name, path))
    return output, date1904


def _xlsx_direct_value(
    cell: ET.Element,
    shared_strings: list[str],
    date_styles: set[int],
    date1904: bool,
) -> Any:
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    cell_type = str(cell.attrib.get("t") or "")
    try:
        style_index = int(cell.attrib.get("s", "0"))
    except Exception:
        style_index = 0

    if cell_type == "inlineStr":
        inline = cell.find(f"{ns}is")
        if inline is None:
            return None
        return "".join(
            node.text or ""
            for node in inline.iter(f"{ns}t")
        )

    value_node = cell.find(f"{ns}v")
    if value_node is None or value_node.text is None:
        return None
    raw_value = value_node.text

    if cell_type == "s":
        try:
            return shared_strings[int(raw_value)]
        except Exception:
            return raw_value
    if cell_type in {"str", "e"}:
        return raw_value
    if cell_type == "b":
        return raw_value == "1"
    if cell_type == "d":
        try:
            return datetime.fromisoformat(raw_value.replace("Z", "+00:00"))
        except Exception:
            return raw_value

    try:
        number = float(raw_value)
    except Exception:
        return raw_value

    if style_index in date_styles:
        base = datetime(1904, 1, 1) if date1904 else datetime(1899, 12, 30)
        try:
            return base + timedelta(days=number)
        except Exception:
            pass

    if number.is_integer():
        return int(number)
    return number


def _read_xlsx_direct(raw: bytes) -> list[dict[str, Any]]:
    """Fallback XML puro para XLSX/XLT(X) incompatível com openpyxl.

    Lê apenas os valores armazenados no pacote Office Open XML e, portanto,
    ignora integralmente formatações, validações e objetos visuais que podem
    conter intervalos inválidos. Fórmulas usam o valor em cache do próprio
    arquivo, equivalente ao comportamento data_only=True.
    """
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    sheets: list[dict[str, Any]] = []

    with zipfile.ZipFile(io.BytesIO(raw), "r") as archive:
        shared_strings = _xlsx_shared_strings(archive)
        date_styles = _xlsx_date_styles(archive)
        manifest, date1904 = _xlsx_workbook_manifest(archive)

        for index, (sheet_name, sheet_path) in enumerate(manifest):
            rows: list[list[Any]] = []
            started = False
            blank_streak = 0
            last_row_number = 0
            max_rows = 250_000
            stop_sheet = False

            stream = io.BytesIO(archive.read(sheet_path))
            for _, element in ET.iterparse(stream, events=("end",)):
                if element.tag != f"{ns}row":
                    continue

                try:
                    row_number = int(
                        element.attrib.get("r") or (last_row_number + 1)
                    )
                except Exception:
                    row_number = last_row_number + 1

                if row_number > max_rows:
                    element.clear()
                    break

                gap = max(row_number - last_row_number - 1, 0)
                if gap:
                    if started and blank_streak + gap >= 150:
                        stop_sheet = True
                    elif started:
                        rows.extend([[] for _ in range(gap)])
                        blank_streak += gap
                    elif row_number <= 51:
                        rows.extend(
                            [[] for _ in range(min(gap, 50 - last_row_number))]
                        )

                if stop_sheet:
                    element.clear()
                    break

                row_values: list[Any] = []
                for cell in element.findall(f"{ns}c"):
                    column_index = _xlsx_column_index(
                        str(cell.attrib.get("r") or "")
                    )
                    if column_index <= 0 or column_index > 256:
                        continue
                    while len(row_values) < column_index:
                        row_values.append(None)
                    row_values[column_index - 1] = _xlsx_direct_value(
                        cell,
                        shared_strings,
                        date_styles,
                        date1904,
                    )

                clean = _trim_workbook_row(tuple(row_values))
                is_blank = not any(
                    value is not None
                    and (
                        not isinstance(value, str)
                        or value.strip()
                    )
                    for value in clean
                )

                if is_blank:
                    if not started:
                        if row_number <= 50:
                            rows.append([])
                    else:
                        blank_streak += 1
                        if blank_streak >= 150:
                            element.clear()
                            break
                        rows.append([])
                else:
                    started = True
                    blank_streak = 0
                    rows.append(clean)

                last_row_number = row_number
                element.clear()

            while rows and not rows[-1]:
                rows.pop()

            sheets.append(
                {
                    "index": index,
                    "name": str(sheet_name),
                    "rows": rows,
                    "row_count": int(len(rows)),
                    "column_count": int(
                        max((len(row) for row in rows), default=0)
                    ),
                }
            )

    return sheets


def _read_xlsx_streaming(raw: bytes) -> list[dict[str, Any]]:
    """Leitura limitada/streaming para relatórios Excel grandes do ERP.

    Evita materializar a área formatada inteira do workbook, causa principal
    dos travamentos durante a normalização no Streamlit.
    """
    try:
        workbook = _load_workbook_streaming(raw)
    except Exception as exc:
        if "MultiCellRange" in str(exc):
            return _read_xlsx_direct(raw)
        raise

    sheets: list[dict[str, Any]] = []
    try:
        for index, sheet_name in enumerate(workbook.sheetnames):
            ws = workbook[sheet_name]
            rows: list[list[Any]] = []
            started = False
            blank_streak = 0
            max_rows = 250_000

            for row_index, values in enumerate(
                ws.iter_rows(values_only=True),
                start=1,
            ):
                if row_index > max_rows:
                    break

                clean = _trim_workbook_row(values)
                is_blank = not any(
                    value is not None
                    and (not isinstance(value, str) or value.strip())
                    for value in clean
                )

                if is_blank:
                    # Preserva linhas iniciais porque alguns relatórios usam
                    # header=1/header=4. Depois que os dados começaram, uma
                    # longa sequência vazia encerra a leitura.
                    if not started:
                        if row_index <= 50:
                            rows.append([])
                        continue
                    blank_streak += 1
                    if blank_streak >= 150:
                        break
                    rows.append([])
                    continue

                started = True
                blank_streak = 0
                rows.append(clean)

            # Remove vazios apenas do fim; nunca do início.
            while rows and not rows[-1]:
                rows.pop()

            sheets.append(
                {
                    "index": index,
                    "name": str(sheet_name),
                    "rows": rows,
                    "row_count": int(len(rows)),
                    "column_count": int(max((len(row) for row in rows), default=0)),
                }
            )
    except Exception as exc:
        if "MultiCellRange" in str(exc):
            return _read_xlsx_direct(raw)
        raise
    finally:
        workbook.close()
    return sheets


def build_normalized_source(
    source_key: str,
    file_name: str,
    raw: bytes,
) -> tuple[bytes, int]:
    """Converte a fonte uma única vez para o pacote técnico SETTA_SOURCE_V1.

    O Excel original continua no Storage para auditoria. Os consumidores
    usam a representação técnica e só recorrem ao bruto como contingência.
    """
    suffix = Path(str(file_name or "").lower()).suffix
    sheets: list[dict[str, Any]] = []

    if suffix in {".xlsx", ".xlsm", ".xltx"}:
        sheets = _read_xlsx_streaming(raw)

    elif suffix == ".xls":
        book = pd.ExcelFile(io.BytesIO(raw), engine="xlrd")
        for index, sheet_name in enumerate(book.sheet_names):
            frame = pd.read_excel(
                book,
                sheet_name=sheet_name,
                header=None,
                dtype=object,
            )
            sheets.append(
                {
                    "index": index,
                    "name": str(sheet_name),
                    "rows": _frame_values(frame),
                    "row_count": int(len(frame)),
                    "column_count": int(frame.shape[1]),
                }
            )

    elif suffix == ".csv":
        frame = None
        last_error: Exception | None = None
        for sep in (None, ";", ",", "\t"):
            try:
                frame = pd.read_csv(
                    io.BytesIO(raw),
                    header=None,
                    dtype=object,
                    sep=sep,
                    engine="python" if sep is None else "c",
                )
                if frame.shape[1] > 1 or sep == "\t":
                    break
            except Exception as exc:
                last_error = exc
                frame = None
        if frame is None:
            raise ValueError(f"Não foi possível normalizar o CSV: {last_error}")
        sheets.append(
            {
                "index": 0,
                "name": "CSV",
                "rows": _frame_values(frame),
                "row_count": int(len(frame)),
                "column_count": int(frame.shape[1]),
            }
        )
    else:
        return b"", 0

    if str(source_key).strip().lower() == "movimentacao":
        if not sheets or not (sheets[0].get("rows") or []):
            raise ValueError("MOVIMENTAÇÃO vazia.")

        header = list((sheets[0].get("rows") or [])[0] or [])
        if len(header) < 12:
            raise ValueError(
                "MOVIMENTAÇÃO fora do padrão: esperado no mínimo A:L, "
                "com C=TM, J=EMISSAO, K=USUARIO e L=ARMAZEM."
            )

        def _h(value):
            text = unicodedata.normalize("NFKD", str(value or ""))
            text = "".join(ch for ch in text if not unicodedata.combining(ch))
            return re.sub(r"[^A-Z0-9]+", "", text.upper())

        esperado = {
            "C": (2, {"TM", "TESA", "TES"}),
            "J": (9, {"EMISSAO", "DATAEMISSAO"}),
            "K": (10, {"USUARIO"}),
            "L": (11, {"ARMAZEM", "ARMZ"}),
        }
        erros = []
        for letra, (idx, aceitos) in esperado.items():
            recebido = _h(header[idx] if idx < len(header) else "")
            if recebido not in aceitos:
                erros.append(
                    f"{letra} esperado {sorted(aceitos)[0]} e veio "
                    f"'{header[idx] if idx < len(header) else ''}'"
                )

        if erros:
            raise ValueError(
                "MOVIMENTAÇÃO incompatível com o relatório padrão do Protheus: "
                + "; ".join(erros)
            )

    payload = {
        "format": "SETTA_SOURCE_V1",
        "source_key": str(source_key),
        "file_name": str(file_name),
        "sheets": sheets,
    }
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    normalized = gzip.compress(encoded, compresslevel=6)
    return normalized, sum(int(sheet["row_count"]) for sheet in sheets)



def _publish_normalized_source(
    source_key: str,
    normalized: bytes,
    normalized_rows: int,
    source_version: int,
) -> dict:
    prepared = central_api_call(
        "source_normalized_upload_prepare",
        {"source_key": source_key},
        timeout=45,
    )
    path = str(prepared.get("path") or "")
    token = str(prepared.get("token") or "")
    if not path or not token:
        raise RuntimeError("A Central não autorizou o upload da fonte normalizada.")

    client = central_client()
    client.storage.from_("setta-data").upload_to_signed_url(
        path=path,
        token=token,
        file=normalized,
    )

    return central_api_call(
        "source_normalized_commit",
        {
            "source_key": source_key,
            "source_version": int(source_version),
            "rows_count": int(normalized_rows),
            "format": "SETTA_SOURCE_V1",
            "profile": "WORKBOOK_MATRIX_V1",
            "content_sha256": hashlib.sha256(normalized).hexdigest(),
        },
        timeout=45,
    ).get("data") or {}


def backfill_normalized_source(source_key: str) -> dict:
    """Normaliza uma fonte já existente sem alterar sua versão de negócio."""
    meta = central_api_call(
        "source_download",
        {"source_key": source_key},
        timeout=30,
    ).get("data") or {}
    signed_url = str(meta.get("signed_url") or "")
    if not signed_url:
        raise RuntimeError(f"Fonte {source_key} sem arquivo bruto disponível.")

    response = requests.get(signed_url, timeout=180)
    response.raise_for_status()
    raw = response.content
    file_name = str(meta.get("last_file_name") or source_key)

    normalized, normalized_rows = build_normalized_source(
        source_key,
        file_name,
        raw,
    )
    if not normalized:
        raise ValueError(
            f"A fonte {source_key} não está em um formato tabular normalizável."
        )

    return _publish_normalized_source(
        source_key,
        normalized,
        normalized_rows,
        int(meta.get("version") or 0),
    )


def save_report(
    source_key: str,
    file_name: str,
    raw: bytes,
    rows_count: int = 0,
    origin: str = "UPLOAD",
    max_attempts: int = 3,
) -> dict:
    """Publica o bruto primeiro e normaliza depois.

    Isso evita que uma planilha grande bloqueie a atualização da fonte antes
    de o arquivo original ser salvo. Se a normalização falhar, o bruto fica
    registrado e uma nova tentativa faz apenas o backfill técnico.
    """
    attempts = max(1, int(max_attempts or 1))
    last_error: Exception | None = None
    committed: dict[str, Any] | None = None

    # MOVIMENTAÇÃO possui layout operacional fixo. Valida antes de substituir
    # a fonte corrente para impedir que um relatório tratado/incompatível
    # sobrescreva o relatório padrão usado pelos indicadores.
    pre_normalized: bytes | None = None
    pre_normalized_rows = 0
    if str(source_key).strip().lower() == "movimentacao":
        pre_normalized, pre_normalized_rows = build_normalized_source(
            source_key,
            file_name,
            raw,
        )
        if not pre_normalized:
            raise ValueError("MOVIMENTAÇÃO não pôde ser normalizada.")

    # 1) Salva imediatamente o original e registra a nova versão.
    for attempt in range(1, attempts + 1):
        try:
            prepared = central_api_call(
                "source_upload_prepare",
                {"source_key": source_key},
                timeout=45,
            )
            path = str(prepared.get("path") or "")
            token = str(prepared.get("token") or "")
            if not path or not token:
                raise RuntimeError("A Central não retornou autorização para upload.")

            client = central_client()
            client.storage.from_("setta-data").upload_to_signed_url(
                path=path,
                token=token,
                file=raw,
            )

            content_sha256 = hashlib.sha256(raw).hexdigest()
            committed = central_api_call(
                "source_commit",
                {
                    "source_key": source_key,
                    "file_name": file_name,
                    "rows_count": int(rows_count),
                    "mime_type": "application/octet-stream",
                    "content_sha256": content_sha256,
                },
                timeout=45,
            ).get("data") or {}
            break
        except Exception as exc:
            last_error = exc
            if attempt < attempts:
                time.sleep(1.5 * attempt)

    if not committed:
        raise RuntimeError(
            f"Falha ao publicar {source_key} após {attempts} tentativa(s): {last_error}"
        ) from last_error

    # 2) Converte uma única vez usando leitura streaming/limitada.
    try:
        if pre_normalized is not None:
            normalized, normalized_rows = pre_normalized, pre_normalized_rows
        else:
            normalized, normalized_rows = build_normalized_source(
                source_key,
                file_name,
                raw,
            )
        if not normalized:
            raise ValueError(
                f"A fonte {source_key} precisa ser tabular para entrar na Central normalizada."
            )
    except Exception as exc:
        raise RuntimeError(
            "Arquivo bruto salvo e versão atualizada, mas a normalização "
            f"técnica não foi concluída: {exc}"
        ) from exc

    normalized_error: Exception | None = None
    for attempt in range(1, attempts + 1):
        try:
            normalized_meta = _publish_normalized_source(
                source_key,
                normalized,
                normalized_rows,
                int(committed.get("version") or 0),
            )
            return {**committed, "normalized": normalized_meta}
        except Exception as exc:
            normalized_error = exc
            if attempt < attempts:
                time.sleep(1.5 * attempt)

    raise RuntimeError(
        "Arquivo bruto salvo e versão atualizada, mas a normalização "
        f"técnica ficou pendente: {normalized_error}"
    ) from normalized_error


def touch_source(source_key: str) -> dict:
    """Registra uma nova conferência do arquivo sem criar nova versão."""
    return central_api_call(
        "source_touch",
        {"source_key": str(source_key)},
        timeout=30,
    ).get("data") or {}



def list_imports(source_key: str | None = None, limit: int = 100) -> list[dict]:
    client = get_client()
    if client:
        query = client.table("data_imports").select("*")
        if source_key:
            query = query.eq("source_key", source_key)
        return query.order("imported_at", desc=True).limit(limit).execute().data or []
    rows = _load_local()["imports"]
    if source_key:
        rows = [x for x in rows if x.get("source_key") == source_key]
    return sorted(rows, key=lambda x: x.get("imported_at", ""), reverse=True)[:limit]


def load_visual_config() -> dict:
    try:
        row = central_api_call(
            "visual_get",
            {"app_key": "setta_global"},
            timeout=30,
        ).get("data") or {}
        return {
            "logo_data": row.get("logo_data") or "",
            "logo_mime": row.get("logo_mime") or "image/png",
            "favicon_data": row.get("favicon_data") or "",
            "favicon_mime": row.get("favicon_mime") or "image/png",
        }
    except Exception:
        return dict(_load_local().get("visual_config") or {})


def save_visual_config(
    logo_data: str | None = None,
    logo_mime: str | None = None,
    favicon_data: str | None = None,
    favicon_mime: str | None = None,
) -> dict:
    current = load_visual_config()
    if logo_data is not None:
        current["logo_data"] = logo_data
    if logo_mime is not None:
        current["logo_mime"] = logo_mime
    if favicon_data is not None:
        current["favicon_data"] = favicon_data
    if favicon_mime is not None:
        current["favicon_mime"] = favicon_mime

    try:
        row = central_api_call(
            "visual_set",
            {
                "app_key": "setta_global",
                **current,
            },
            timeout=30,
        ).get("data") or current
        return {
            "logo_data": row.get("logo_data") or "",
            "logo_mime": row.get("logo_mime") or "image/png",
            "favicon_data": row.get("favicon_data") or "",
            "favicon_mime": row.get("favicon_mime") or "image/png",
        }
    except Exception:
        data = _load_local()
        data["visual_config"] = current
        _save_local(data)
        return current


def reset_visual_config() -> None:
    save_visual_config(
        logo_data="",
        logo_mime="image/png",
        favicon_data="",
        favicon_mime="image/png",
    )
